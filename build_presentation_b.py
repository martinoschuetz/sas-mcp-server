import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def build_presentation_b():
    template_path = r'C:\Git\martinoschuetz\projects\CandF\SAS External  PowerPoint.pptx'
    output_path = r'C:\Git\sas-mcp-server\Presentation_B_Autonomous_Analytics_SAS_Viya.pptx'

    prs = pptx.Presentation(template_path)

    # Clear existing template slides while preserving slide masters, layouts, and themes
    for i in range(len(prs.slides) - 1, -1, -1):
        rId = prs.slides._sldIdLst[i].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[i]

    layout_title = prs.slide_layouts[0]     # SAS - Title
    layout_content = prs.slide_layouts[3]   # SAS - Title & Subtitle
    layout_closing = prs.slide_layouts[8]   # SAS - Closing

    # SAS Brand Color Palette (Official Guidelines)
    c_navy = RGBColor(1, 32, 54)        # #012036 Slate/Navy
    c_blue = RGBColor(51, 163, 255)     # #33A3FF SAS Blue
    c_teal = RGBColor(41, 214, 205)     # #29D6CD Teal
    c_green = RGBColor(21, 181, 123)    # #15B57B Green
    c_purple = RGBColor(109, 105, 255)  # #6D69FF Purple
    c_amber = RGBColor(255, 204, 50)    # #FFCC32 Amber/Gold
    c_wine = RGBColor(134, 19, 79)      # #86134F Wine/Burgundy
    c_card_bg = RGBColor(250, 252, 254) # Off-white card fill
    c_card_border = RGBColor(215, 225, 235)
    c_gray = RGBColor(85, 100, 115)     # Subdued gray text
    c_dark = RGBColor(20, 30, 42)       # Primary dark body text

    font_family = "Anova Light"

    def set_slide_header(slide, title_text, subtitle_text):
        for sh in slide.shapes:
            if sh.is_placeholder:
                if sh.placeholder_format.idx == 0:
                    sh.text_frame.word_wrap = True
                    p = sh.text_frame.paragraphs[0]
                    p.text = title_text
                    p.font.name = font_family
                    p.font.bold = True
                    p.font.size = Pt(36)
                    p.font.color.rgb = c_navy
                elif sh.placeholder_format.idx == 10:
                    sh.text_frame.word_wrap = True
                    p = sh.text_frame.paragraphs[0]
                    p.text = subtitle_text
                    p.font.name = font_family
                    p.font.size = Pt(17)
                    p.font.color.rgb = c_gray

    def set_slide_notes(slide, notes_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text.strip()

    def add_card(slide, left, top, width, height, accent_color, header_title, items, subtitle=None, mode='wide'):
        # Container shape
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = c_card_bg
        card.line.color.rgb = c_card_border
        card.line.width = Pt(1.2)

        # Top Accent Ribbon
        ribbon_h = Inches(0.12)
        ribbon = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, ribbon_h)
        ribbon.fill.solid()
        ribbon.fill.fore_color.rgb = accent_color
        ribbon.line.fill.background()

        # Text Frame
        if mode == 'wide': # 3 cards per slide
            pad_x = Inches(0.24)
            pad_y = Inches(0.24)
            title_size = Pt(21)
            sub_size = Pt(14)
            item_head_size = Pt(14.5)
            item_body_size = Pt(14)
            space_after = Pt(13)
            sub_space = Pt(14)
        elif mode == 'medium': # 4 cards per slide
            pad_x = Inches(0.20)
            pad_y = Inches(0.20)
            title_size = Pt(18.5)
            sub_size = Pt(13)
            item_head_size = Pt(13.5)
            item_body_size = Pt(13)
            space_after = Pt(9)
            sub_space = Pt(11)
        else: # compact, 5 cards per slide
            pad_x = Inches(0.18)
            pad_y = Inches(0.22)
            title_size = Pt(18)
            sub_size = Pt(12.5)
            item_head_size = Pt(13)
            item_body_size = Pt(12.5)
            space_after = Pt(8)
            sub_space = Pt(10)

        tb = slide.shapes.add_textbox(left + pad_x, top + pad_y, width - 2 * pad_x, height - 2 * pad_y)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        # Card Title
        p_title = tf.paragraphs[0]
        p_title.text = header_title
        p_title.font.name = font_family
        p_title.font.bold = True
        p_title.font.size = title_size
        p_title.font.color.rgb = c_navy
        p_title.space_after = Pt(4)

        if subtitle:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle
            p_sub.font.name = font_family
            p_sub.font.size = sub_size
            p_sub.font.bold = False
            p_sub.font.color.rgb = accent_color
            p_sub.space_after = sub_space
        else:
            p_title.space_after = Pt(16) if mode == 'wide' else Pt(12)

        for heading, body in items:
            p_item = tf.add_paragraph()
            p_item.space_after = space_after

            run_head = p_item.add_run()
            run_head.text = f"•  {heading}: " if heading else "•  "
            run_head.font.name = font_family
            run_head.font.bold = True
            run_head.font.size = item_head_size
            run_head.font.color.rgb = c_navy

            run_body = p_item.add_run()
            run_body.text = body
            run_body.font.name = font_family
            run_body.font.bold = False
            run_body.font.size = item_body_size
            run_body.font.color.rgb = c_dark

        return card

    # Standard coordinates for 3-card layout (Width: 20.0 inches)
    margin_x_3 = Inches(1.38)
    card_y = Inches(2.55)
    card_w_3 = Inches(5.42)
    card_h = Inches(7.30)
    gap_3 = Inches(0.48)

    # Standard coordinates for 4-card layout (Width: 20.0 inches)
    margin_x_4 = Inches(1.00)
    card_w_4 = Inches(4.20)
    gap_4 = Inches(0.40)

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(layout_title)
    for sh in s1.shapes:
        if sh.is_placeholder:
            if sh.placeholder_format.idx == 0:
                p = sh.text_frame.paragraphs[0]
                p.text = "Autonomous Industrial Analytics on SAS Viya"
                p.font.name = font_family
                p.font.bold = True
                p.font.size = Pt(44)
            elif sh.placeholder_format.idx == 10:
                p = sh.text_frame.paragraphs[0]
                p.text = "Bridging Real-Time Streaming, Custom IoT Frameworks, Quality Analytics, and Forecasting via Model Context Protocol (MCP)"
                p.font.name = font_family
                p.font.size = Pt(19)
            elif sh.placeholder_format.idx == 11:
                p = sh.text_frame.paragraphs[0]
                p.text = "SAS Viya Architecture & Applied AI Showcase | 30–45 Minute Overview"
                p.font.name = font_family
                p.font.size = Pt(15)

    set_slide_notes(
        s1,
        "Welcome everyone. Today we are presenting a comprehensive architectural and practical overview of autonomous industrial analytics on SAS Viya.\n\n"
        "Over the past year, we have tackled a critical challenge: How do you bridge modern Large Language Model agents with heavy, deterministic enterprise analytics? "
        "Too often, AI demos stop at toy chatbots generating unverified python scripts. In industrial manufacturing, pharmaceutical fermentation, warranty quality, and supply chain forecasting, "
        "you cannot afford hallucinations, and you cannot afford unverified numbers.\n\n"
        "This presentation covers how we used the open Model Context Protocol (MCP) to turn SAS Viya into an agent-driven analytical engine across four major domains: "
        "Event Stream Processing (ESP), the Custom Analysis Framework (CAF) for IoT, SAS Field Quality Analytics (FQA), and Distributed Hierarchical Forecasting with ATSM. "
        "We will walk through the business motivation, the production architecture, the code-level implementations, and how this runs in live environments."
    )

    # =========================================================================
    # SLIDE 2: Motivation & The Enterprise AI Dilemma
    # =========================================================================
    s2 = prs.slides.add_slide(layout_content)
    set_slide_header(
        s2,
        "The Enterprise AI Dilemma: Hallucinations vs. Industrial Rigor",
        "Why Generic LLMs Fail in Heavy Industry, and How MCP Transforms SAS Viya into a Deterministic AI Engine"
    )

    add_card(
        s2, margin_x_3, card_y, card_w_3, card_h, c_wine,
        "The LLM Reality & Failure Modes",
        [
            ("Statistical Hallucination", "LLMs generate plausible-looking numbers and fabricated correlations when asked to perform numerical or statistical reasoning."),
            ("No Execution Engine", "Frontier models cannot natively query massive in-memory CAS tables, trigger SAS PROCs, or monitor live Kafka streaming pipelines."),
            ("State & Context Amnesia", "Chatbots lack long-term memory of industrial data models, hierarchical taxonomies, and multi-step investigation histories."),
            ("Security & Governance", "Sending raw proprietary industrial telemetry or unredacted warranty claim text directly to public AI endpoints violates enterprise compliance.")
        ],
        subtitle="Unassisted Frontier Models",
        mode='wide'
    )

    add_card(
        s2, margin_x_3 + card_w_3 + gap_3, card_y, card_w_3, card_h, c_blue,
        "The SAS Viya Foundation",
        [
            ("Deterministic Analytics", "SAS Cloud Analytic Services (CAS) processes billions of rows in memory with mathematically proven algorithms."),
            ("Industrial Streaming", "SAS Event Stream Processing (ESP) executes complex event processing, sliding windows, and scoring at microsecond latency."),
            ("Domain Engines", "Specialized solutions for Field Quality Analytics (FQA), IoT custom frameworks (CAF), and Automatic Time Series Modeling (ATSM)."),
            ("Enterprise Trust", "Auditable logs, governed data access, role-based security, and predictable enterprise deployment on Kubernetes.")
        ],
        subtitle="The Analytical Workhorse",
        mode='wide'
    )

    add_card(
        s2, margin_x_3 + 2 * (card_w_3 + gap_3), card_y, card_w_3, card_h, c_green,
        "The MCP Bridge",
        [
            ("Open Standard Protocol", "Model Context Protocol (MCP) standardizes how AI agents discover tools, inspect schemas, and execute remote capabilities."),
            ("Numbers from SAS, Words from AI", "The agent reasons, plans, and synthesizes, but delegates every statistical calculation to verified SAS execution engines."),
            ("Bidirectional Telemetry", "Tools return structured JSON, execution logs, and embedded HTML micro-views directly into the agent reasoning context."),
            ("Zero-Click Automation", "Transforms manual GUI workflows (Model Studio, ESP Studio, FQA) into headless, programmatic execution pipelines.")
        ],
        subtitle="The Agent-to-Engine Interface",
        mode='wide'
    )

    set_slide_notes(
        s2,
        "Slide 2 sets the foundation for our entire approach.\n\n"
        "On the left, we look honestly at LLMs today. They have incredible general reasoning, coding, and synthesis capabilities. But in high-stakes environments—like automotive safety recalls or biologics fermentation—"
        "an LLM cannot be trusted to do mental math or invent statistical metrics. If an LLM tells you that a warranty failure rate increased by 23%, but hallucinated the denominator, you have an engineering crisis.\n\n"
        "In the center, SAS Viya represents the gold standard of deterministic computation. CAS handles terabytes in memory, ESP crunches thousands of sensor events per second, and ATSM runs proven time series algorithms.\n\n"
        "On the right is the breakthrough: Model Context Protocol. By exposing SAS Viya capabilities through typed, governed MCP tool contracts, we give AI agents the ability to operate SAS as their analytical brain. "
        "Our golden rule is: 'Numbers come from SAS, words come from you.'"
    )

    # =========================================================================
    # SLIDE 3: MCP Server Architecture & SAS RAM
    # =========================================================================
    s3 = prs.slides.add_slide(layout_content)
    set_slide_header(
        s3,
        "SAS Viya MCP Server: Modular 15-Tier Architecture",
        "Enterprise Concurrency, Containerized Deployment on SAS RAM, and High-Throughput REST/CAS Gateways"
    )

    add_card(
        s3, margin_x_3, card_y, card_w_3, card_h, c_blue,
        "Protocol & Async Concurrency",
        [
            ("FastMCP Framework", "High-performance Python 3.12 architecture utilizing anyio structured concurrency for non-blocking I/O operations."),
            ("Dual Transports", "Supports both standard stdio (for local CLI / IDE agents) and Server-Sent Events (SSE / HTTP) for networked enterprise platforms."),
            ("Configurable Timeouts", "Tuned via VIYA_CLIENT_TIMEOUT (set to 1800s in our production image) to allow long-running CAS modeling and batch forecasting."),
            ("App UI Extensions", "Emits interactive sandboxed HTML views directly alongside JSON results for tabular data, SAS logs, and visual reports.")
        ],
        subtitle="Core Engine & Runtime",
        mode='wide'
    )

    add_card(
        s3, margin_x_3 + card_w_3 + gap_3, card_y, card_w_3, card_h, c_teal,
        "Modular 15-Tier Tool System",
        [
            ("Foundation Tiers (0–4)", "Session management, CAS data exploration, dynamic PROC execution, and batch job status monitoring."),
            ("Analytics Tiers (5–8)", "Automated Machine Learning (AutoML), Visual Analytics report authoring, Model Scoring (MAS), and Decisioning flows."),
            ("Domain Tiers (9–13)", "Field Quality Analytics (Tier 10), Event Stream Processing (Tier 12), Custom Analysis Framework (Tier 13)."),
            ("Knowledge RAG (Tier 14)", "Self-disclosing skill and documentation system allowing agents to dynamically query specialized domain playbooks.")
        ],
        subtitle="Tier 0 through Tier 14",
        mode='wide'
    )

    add_card(
        s3, margin_x_3 + 2 * (card_w_3 + gap_3), card_y, card_w_3, card_h, c_purple,
        "Containerization & SAS RAM",
        [
            ("SAS RAM Packaging", "Built and published to sasram.azurecr.io/sas-mcp-server-msz:v.1.18.1 for enterprise agent orchestration."),
            ("Hardened Docker Build", "Multi-stage build utilizing uv and pinned dependencies with strict secret isolation (.dockerignore & .gitignore hardening)."),
            ("Viya Auth Integration", "Dynamic OAuth2 Bearer token exchange with automatic token refresh and encrypted session caching."),
            ("RBAC & Safety Modes", "Configurable read-only tool gating (_access.py) ensuring agents cannot modify production assets without explicit permission.")
        ],
        subtitle="SAS Retrieval Agent Manager",
        mode='wide'
    )

    set_slide_notes(
        s3,
        "Now let's examine the technical architecture of the SAS MCP Server.\n\n"
        "The server is implemented in modern Python 3.12 on top of FastMCP. We use structured concurrency with anyio, meaning the server can handle multiple concurrent agent requests without blocking.\n\n"
        "Notice the timeout handling: standard HTTP timeouts of 30 or 60 seconds fail immediately when running large CAS analytical jobs or complex ATSM forecasting tournaments. "
        "We engineered the server with VIYA_CLIENT_TIMEOUT set to 1800 seconds (30 minutes), enabling heavy industrial workflows.\n\n"
        "The tools are partitioned into 15 logical tiers. An agent working on warranty only loads Tier 10 tools; an agent authoring IoT analyses loads Tier 13. "
        "Furthermore, we package the server into a hardened Docker container published to Azure Container Registry for SAS Retrieval Agent Manager (RAM), allowing seamless deployment across enterprise AI platforms."
    )

    # =========================================================================
    # SLIDE 4: Core Guardrails & The Agent Reliability Contract
    # =========================================================================
    s4 = prs.slides.add_slide(layout_content)
    set_slide_header(
        s4,
        "Deterministic Reliability: The Agent Governance Contract",
        "Rigorous Statistical Guardrails, Preflight Checks, and Zero-Hallucination Execution Principles"
    )

    add_card(
        s4, margin_x_3, card_y, card_w_3, card_h, c_green,
        "Strict Statistical Discipline",
        [
            ("No Mental Arithmetic", "The agent is explicitly forbidden from computing rates, ratios, or correlations in its head; every number must originate in SAS."),
            ("Audited Metric Lineage", "All claims cited in agent executive summaries include direct references to CAS table names, column hashes, and procedure outputs."),
            ("Count Units, Not Claims", "Enforces unit-level counting (PRODUCT_ID) rather than inflated claim pairs, preventing multi-fold statistical bias."),
            ("Correlation vs. Causation", "Rules require at least three supporting findings across two separate analysis types before claiming causal relationships.")
        ],
        subtitle="Mathematical Integrity",
        mode='wide'
    )

    add_card(
        s4, margin_x_3 + card_w_3 + gap_3, card_y, card_w_3, card_h, c_amber,
        "Preflight Validation Checks",
        [
            ("Zero-Variance Detection", "Verifies target variables have sufficient variance before triggering expensive decision trees or regression actions."),
            ("Minimum Event Thresholds", "Blocks execution of statistical models if the filtered cohort has fewer than the required event threshold (e.g., n < 30)."),
            ("High-Cardinality Safeguards", "Analyzes categorical predictors prior to modeling; automatically flags or groups variables exceeding memory bounds."),
            ("Cluster Protection", "Prevents Kubernetes pod OOM crashes by rejecting unviable analytical configurations before compute dispatch.")
        ],
        subtitle="Computational Feasibility",
        mode='wide'
    )

    add_card(
        s4, margin_x_3 + 2 * (card_w_3 + gap_3), card_y, card_w_3, card_h, c_wine,
        "Security & Workspace Hygiene",
        [
            ("Automated PII Redaction", "Technician claim comments and free-text notes are stripped of personal names, vehicle VINs, and locations before context ingest."),
            ("Zero Ghost Nodes", "Guarantees that failed intermediate runs or temporary data selections are cleanly deleted, leaving the customer's UI pristine."),
            ("Untrusted Text Protocol", "Free-text comments are treated strictly as unverified data, preventing prompt-injection attacks disguised as warranty feedback."),
            ("Role-Based Read/Write", "Write operations (creating projects, promoting CAS tables, updating models) require explicit human-in-the-loop authorization.")
        ],
        subtitle="Governance & Privacy",
        mode='wide'
    )

    set_slide_notes(
        s4,
        "Slide 4 addresses the most critical question any enterprise architect will ask: 'How do you ensure the agent doesn't do something crazy?'\n\n"
        "We codified a strict governance contract into the agent's system prompt and skill definitions:\n"
        "First, 'Numbers come from SAS, words come from you.' If an agent is investigating a warranty claim or a bioreactor sensor drift, it cannot summarize numbers it didn't pull directly from a SAS output table.\n\n"
        "Second, Preflight checks. In traditional UI environments, an analyst might configure a decision tree on a dataset with zero events or 50,000 distinct categorical levels, causing the compute pod to grind to a halt or crash. "
        "Our agent runs programmatic preflights: it measures event counts and cardinality first. If the data isn't viable, it tells the user immediately without firing the heavy CAS action.\n\n"
        "Third, workspace hygiene and security. The agent automatically redacts PII from technician comments and cleans up all temporary CAS tables and data selection nodes."
    )

    # =========================================================================
    # SLIDE 5: Domain 1 (ESP) - Motivation & Use Case: Bioreactor Telemetry
    # =========================================================================
    s5 = prs.slides.add_slide(layout_content)
    set_slide_header(
        s5,
        "Streaming IoT: Continuous Bioreactor Telemetry (Celivion C&F)",
        "High-Frequency IoT Event Processing for Fed-Batch Fermentation, Yield Forecasting & Process Safety"
    )

    add_card(
        s5, margin_x_3, card_y, card_w_3, card_h, c_blue,
        "Biological Process Scale",
        [
            ("Commercial Bioreactor Train", "Continuous monitoring across three industrial vessels: bio01, bio02, and bio03 (production scale up to 5,000 L)."),
            ("Complex Biological Dynamics", "Slow 10 to 14-day mammalian cell cultivation cycles per commercial batch, requiring staggered scheduling across suites."),
            ("10 Operational Phases", "Vessels transition through Inoculation, Exponential Growth, Nutrient Feeding, Stationary Expression, and Harvest."),
            ("Sterilization & Turnaround", "Post-harvest drain, Clean-in-Place (CIP), Steam-in-Place (SIP) sterilization, and media refill cycles.")
        ],
        subtitle="Physical Process Setup",
        mode='wide'
    )

    add_card(
        s5, margin_x_3 + card_w_3 + gap_3, card_y, card_w_3, card_h, c_teal,
        "Critical Process Variables",
        [
            ("Critical Quality Attributes (CQAs)", "Real-time tracking of Viability (>85%), Viable Cell Density (VCD), Glucose consumption, and Lactate accumulation."),
            ("Critical Process Parameters (CPPs)", "Tight control loops on Broth Temperature (37.0 °C), Dissolved Oxygen (DO), pH titration, and Agitation (RPM)."),
            ("Vessel Aeration & Off-Gas", "Continuous telemetry for sparger air flow rate, headspace vessel pressure (bar), and foam level sensor signals."),
            ("Specification Envelopes", "Rigorous phase-specific upper and lower operating boundaries (documented in spec_limits_bmc.csv).")
        ],
        subtitle="14 Real-Time Sensors",
        mode='wide'
    )

    add_card(
        s5, margin_x_3 + 2 * (card_w_3 + gap_3), card_y, card_w_3, card_h, c_purple,
        "Four High-Value Targets",
        [
            ("Predictive Quality", "Forecasting cell viability and VCD degradation 1 to 12 hours ahead of time to adjust nutrient feeding schedules."),
            ("Predictive Harvest / RUL", "Accurately estimating Remaining Useful Life (RUL) until glucose depletion drops below the critical 2.0 g/L harvest trigger."),
            ("Foam-Out Prevention", "Early detection of foaming conditions ahead of time to trigger anti-foam surfactant dosing and avoid filter clogging."),
            ("Multi-Sensor Anomaly Detection", "Detecting subtle drift, dissolved oxygen collapse, and abnormal metabolic signatures using in-memory SVDD.")
        ],
        subtitle="Analytical Objectives",
        mode='wide'
    )

    set_slide_notes(
        s5,
        "Now let's examine our first major domain: Real-time event streaming with SAS Event Stream Processing (ESP).\n\n"
        "Our case study is with Celivion Therapeutics on the C&F Unified Namespace (UNS) platform, monitoring a train of commercial mammalian cell culture bioreactors.\n\n"
        "In biomanufacturing, living cells produce therapeutic proteins over a 10 to 14-day fermentation cycle. Even brief disruptions in pH, dissolved oxygen, or temperature can lead to batch collapse, costing millions of dollars.\n\n"
        "We are streaming 14 continuous sensor tags across units bio01, bio02, and bio03. Our analytical objectives are fourfold: "
        "Predicting viability drop-offs, forecasting remaining time-to-harvest based on glucose depletion, preventing catastrophic foam-outs, and detecting multivariate anomalies."
    )

    # =========================================================================
    # SLIDE 6: Domain 1 (ESP) - Technical Architecture: Kafka to CAS
    # =========================================================================
    s6 = prs.slides.add_slide(layout_content)
    set_slide_header(
        s6,
        "Streaming Pipeline Architecture: From Kafka UNS to In-Memory ABT",
        "Overcoming Schema Heterogeneity with opaquestring, Python Normalization, and Direct CAS In-Memory Streaming"
    )

    add_card(
        s6, margin_x_4, card_y, card_w_4, card_h, c_blue,
        "1. Kafka UNS Ingestion",
        [
            ("Confluent Cloud Topic", "Ingests uns.dev.uns.clv.gbl.sas with CloudEvents JSON envelopes."),
            ("opaquestring Connector", "w_source window ingests raw JSON string without schema binding, preventing parser crashes on payload variations."),
            ("Autogen Keys & Offsets", "High-throughput blocksize=1000 with disableespmetaconsumer=true for reliable replaying from earliest offsets.")
        ],
        subtitle="w_source Window",
        mode='medium'
    )

    add_card(
        s6, margin_x_4 + card_w_4 + gap_4, card_y, card_w_4, card_h, c_teal,
        "2. Microsecond Flattening",
        [
            ("Embedded Python Window", "w_flatten utilizes embedded Python create() function to normalize nested JSON envelopes into 18 typed columns."),
            ("Microsecond Precision", "Parses ISO-8601 sensor timestamps into microsecond epoch integers for sub-second synchronization."),
            ("Edge Quality Filtering", "Filters out invalid sensor readings (tag_quality='False') directly in the stream before landing.")
        ],
        subtitle="w_flatten Window",
        mode='medium'
    )

    add_card(
        s6, margin_x_4 + 2 * (card_w_4 + gap_4), card_y, card_w_4, card_h, c_amber,
        "3. In-Memory CAS Streaming",
        [
            ("dfesp_cas_adapter", "w_to_cas connector streams normalized events directly into CAS in-memory table PUBLIC.BMC_RAW."),
            ("Micro-Batch Commits", "commitblocks=10 appends events continuously with zero disk thrashing."),
            ("Persisted Sizing", "Over 90,000 continuous sensor readings landed and backed up to BMC_RAW.sashdat with 100% data integrity.")
        ],
        subtitle="w_to_cas Adapter",
        mode='medium'
    )

    add_card(
        s6, margin_x_4 + 3 * (card_w_4 + gap_4), card_y, card_w_4, card_h, c_green,
        "4. Stateful Batch ABT",
        [
            ("Dynamic Batch Detection", "SAS DATA step detects re-inoculation glucose jumps (0 to 8-12 g/L) to assign true sequential Batch_IDs."),
            ("LOCF Imputation", "Last Observation Carried Forward aligns asynchronous sensor bursts into unified 1-minute time slices."),
            ("Target Engineering", "Accurately computes Remaining Useful Life against current batch completion rather than global timeline.")
        ],
        subtitle="generate_abt.sas",
        mode='medium'
    )

    set_slide_notes(
        s6,
        "Slide 6 shows the exact technical pipeline we built in SAS ESP and CAS.\n\n"
        "First, Kafka ingestion: Confluent Cloud streams telemetry as JSON CloudEvents. If you use a rigid ESP XML schema, a single extra field from an upstream sensor will crash the parser. "
        "We used the opaquestring pattern in the w_source window, treating the payload as a raw string and delegating parsing to an embedded Python window (w_flatten).\n\n"
        "Second, the Python window extracts all 14 sensor tags, converts timestamps to microsecond precision, and validates quality flags.\n\n"
        "Third, the w_to_cas adapter pipes data directly into CAS memory without touching disk, landing in PUBLIC.BMC_RAW.\n\n"
        "Fourth, the analytical base table (ABT) creation. Bioreactor sensors publish asynchronously. We wrote a SAS DATA step that performs Last Observation Carried Forward (LOCF) "
        "and uses a glucose spike rule (jumping from near 0 to >8 g/L) to dynamically detect batch boundaries. This fixed a critical issue where time-to-harvest models previously measured against the end of the entire dataset instead of the active batch."
    )

    # =========================================================================
    # SLIDE 7: Domain 2 (CAF) - Motivation: The AIoT Extensibility Gap
    # =========================================================================
    s7 = prs.slides.add_slide(layout_content)
    set_slide_header(
        s7,
        "Custom Analysis Framework (CAF): The AIoT Extensibility Gap",
        "Moving from Proprietary GUI Wizards to Declarative 'Analysis-as-Code' for Enterprise IoT Algorithms"
    )

    add_card(
        s7, margin_x_3, card_y, card_w_3, card_h, c_wine,
        "The Traditional AIoT Dilemma",
        [
            ("Fixed Out-of-the-Box Models", "SAS Analytics for IoT ships standard analytics (Pareto, Reliability, Trend and Control), but enterprise operations require proprietary algorithms."),
            ("The GUI Designer Barrier", "Creating custom analysis types previously required clicking through a multi-step web UI designer—slow, manual, and error-prone."),
            ("No Version Control", "GUI-created analysis models could not be diffed, reviewed in GitHub pull requests, or integrated into modern CI/CD pipelines."),
            ("Unforgiving Silent Failures", "A single misspelled token in the Visual Analytics report definition renders an empty screen without throwing an error.")
        ],
        subtitle="The Proprietary Bottleneck",
        mode='wide'
    )

    add_card(
        s7, margin_x_3 + card_w_3 + gap_3, card_y, card_w_3, card_h, c_blue,
        "Reverse-Engineering the Format",
        [
            ("Decoded 34 Packages", "Systematically reverse-engineered 34 packaged analysis types in CAF_Examples to uncover the true underlying XML and macro schema."),
            ("Undocumented Artifacts", "Documented the exact artifact format for input.xml, SAS CDATA code execution, and Visual Analytics BIRD templates."),
            ("Three Architectural Flavors", "Codified the differences between STANDARD (driver macros), ALERTS (early warning), and CUSTOM (user-authored code)."),
            ("Comprehensive Guide", "Authored CAF_documentation.md: a complete authoring reference for hand-crafting analysis types without the UI.")
        ],
        subtitle="Analysis-as-Code Breakthrough",
        mode='wide'
    )

    add_card(
        s7, margin_x_3 + 2 * (card_w_3 + gap_3), card_y, card_w_3, card_h, c_green,
        "Enterprise Business Impact",
        [
            ("Programmatic Hot-Reloading", "Analysis types can be authored in VS Code, validated against a linter, and deployed instantly via MCP REST APIs."),
            ("Support Vector Data Description", "Implemented SVDD Anomaly Detection: a one-class classification model for complex multi-sensor industrial equipment."),
            ("Reusable Step Templates", "Modular packaging allows data scientists to build an analytical algorithm once and expose it as a self-service tool to plant engineers."),
            ("Automated Report Binding", "Visual Analytics dashboards are parameterized to bind automatically to newly promoted CAS output tables.")
        ],
        subtitle="Agile Analytical Delivery",
        mode='wide'
    )

    set_slide_notes(
        s7,
        "Now let's examine our second major domain: The Custom Analysis Framework (CAF) in SAS Analytics for IoT.\n\n"
        "SAS Analytics for IoT is a premier enterprise solution, but historically it had a major bottleneck: creating a custom analysis type required clicking through a GUI wizard. "
        "You couldn't store an analysis type in Git. You couldn't run a diff on a pull request. And if a parameter name was mistyped, the Visual Analytics report failed silently with a blank screen.\n\n"
        "We undertook a systematic reverse-engineering project, analyzing 34 standard and custom CAF packages. We decoded the exact structure of input.xml, the embedded SAS CDATA execution blocks, and the Visual Analytics BIRD XML templates.\n\n"
        "This transformed CAF into 'Analysis-as-Code'. Data scientists can now write algorithms like SVDD (Support Vector Data Description) in Python or SAS, package them programmatically, and deploy them directly to Viya via the MCP server."
    )

    # =========================================================================
    # SLIDE 8: Domain 2 (CAF) - Technical Architecture & SVDD Model
    # =========================================================================
    s8 = prs.slides.add_slide(layout_content)
    set_slide_header(
        s8,
        "CAF Deep-Dive: Deconstructing the 3-Layer Artifact",
        "Input Form Specification, Server-Side Compute CDATA, and Visual Analytics BIRD XML Templatization"
    )

    add_card(
        s8, margin_x_3, card_y, card_w_3, card_h, c_blue,
        "Layer 1: Input & Form Spec",
        [
            ("input.xml Definition", "Defines the UI form the end user sees: parameter pickers, dropdowns, numeric sliders, and container groupings."),
            ("Dynamic Lookups", "Uses platform lookups with RESOLVEONUI to populate variable dropdowns dynamically from selected CAS tables."),
            ("Interaction Rules", "Declarative conditional logic that shows, hides, or disables parameters based on user selections without JavaScript."),
            ("Embedded SAS Code", "Encapsulates the complete SAS execution driver inside a <Code><![CDATA[...]]></Code> block with macro token substitution.")
        ],
        subtitle="UI & Orchestration Contract",
        mode='wide'
    )

    add_card(
        s8, margin_x_3 + card_w_3 + gap_3, card_y, card_w_3, card_h, c_amber,
        "Layer 2: Compute Execution",
        [
            ("PROC CAS & Action Sets", "Executes high-performance CAS action sets (e.g., svm.svddTrain, svm.svddScore) across distributed worker nodes."),
            ("Automated Variable Resolution", "Resolves user-selected input tags into macro arrays (&varList) and sets hyperparameter values safely."),
            ("In-Memory Promotion", "Persists scored analytical results into promoted CAS tables with consistent naming conventions for downstream reporting."),
            ("Telemetry & Logging", "Emits structured execution logs into the AIoT run repository for full historical auditability.")
        ],
        subtitle="Distributed Server Execution",
        mode='wide'
    )

    add_card(
        s8, margin_x_3 + 2 * (card_w_3 + gap_3), card_y, card_w_3, card_h, c_purple,
        "Layer 3: Visual Analytics BIRD",
        [
            ("vaoutput.template", "Visual Analytics report definition exported as raw BIRD (Business Intelligence Report Definition) XML."),
            ("Tokenized Data Sources", "Replaces hardcoded CAS library and table names with parameterized tokens (e.g., {{OUTPUT_TABLE_NAME}})."),
            ("Rich Pre-Built Visuals", "Automatically configures scatter plots, anomaly score timelines, ROC curves, and outlier threshold bands."),
            ("Zero-Click User Experience", "When the analysis run finishes, the interactive Visual Analytics report opens immediately, fully populated.")
        ],
        subtitle="Interactive VA Reporting",
        mode='wide'
    )

    set_slide_notes(
        s8,
        "Slide 8 breaks down the three architectural layers that make up a Custom Analysis Framework package.\n\n"
        "Layer 1 is input.xml. This is a declarative schema specifying the UI form. It defines parameter containers, validation ranges, cascading lookups, and conditional interaction rules. "
        "Crucially, it also houses the SAS driver code in a CDATA block.\n\n"
        "Layer 2 is Compute. When the user or agent clicks 'Run', the Viya backend substitutes the form parameters into the SAS code and executes it in CAS. "
        "In our SVDD implementation, it calls the svm.svddTrain action set, finding the minimal enclosing hypersphere around normal operating data to detect outliers.\n\n"
        "Layer 3 is the Visual Analytics output template (BIRD XML). Generating BIRD XML by hand is notoriously difficult. "
        "Instead, we design a beautiful report in the Visual Analytics UI, and then use our automated MCP tools to templatize the XML, replacing hardcoded table references with dynamic tokens."
    )

    # =========================================================================
    # SLIDE 9: Domain 2 (CAF) - MCP Authoring Tools
    # =========================================================================
    s9 = prs.slides.add_slide(layout_content)
    set_slide_header(
        s9,
        "Autonomous CAF Authoring: MCP Tools for Code-First IoT",
        "Schema Introspection, Pre-Flight Specification Linting, and Visual Report Reverse-Engineering"
    )

    add_card(
        s9, margin_x_3, card_y, card_w_3, card_h, c_teal,
        "describe_caf_schema",
        [
            ("Progressive Disclosure", "Allows an AI agent to query the CAF XML contract layer-by-layer (parameters, lookups, interactions, code, BIRD XML)."),
            ("Prevents Context Bloat", "Instead of stuffing a 50-page XML manual into the agent's prompt, the tool discloses schemas on demand."),
            ("Display Control Reference", "Details the exact XML properties required for dropdowns, date pickers, numeric inputs, and radio groups."),
            ("Error-Free Authoring", "Equips LLMs with the exact vocabulary required to author valid packages from scratch.")
        ],
        subtitle="Tier 13: Schema Introspection",
        mode='wide'
    )

    add_card(
        s9, margin_x_3 + card_w_3 + gap_3, card_y, card_w_3, card_h, c_amber,
        "validate_analysis_type_spec",
        [
            ("Pre-Upload Static Linter", "Scans user- or agent-generated CAF packages before sending them to the Viya REST API."),
            ("Token Producer Verification", "Validates that every token referenced in the Visual Analytics report is produced by a parameter in input.xml."),
            ("Circular Cascade Detection", "Detects self-retriggering interaction loops and invalid #ctx method calls in the input layer."),
            ("Actionable Error Diagnostics", "Returns precise line numbers and exact remediation instructions instead of generic HTTP 500 errors.")
        ],
        subtitle="Tier 13: Preflight Validation",
        mode='wide'
    )

    add_card(
        s9, margin_x_3 + 2 * (card_w_3 + gap_3), card_y, card_w_3, card_h, c_green,
        "templatize_va_report & Deploy",
        [
            ("Automated BIRD Conversion", "templatize_va_report reads an existing interactive report from the Viya folder service and tokenizes it automatically."),
            ("CI/CD Package Assembly", "build_output_contract zips the input XML, SAS code, and report template into a production-ready CAF bundle."),
            ("create_analysis_type", "Registers the package directly with the /iotAnalysisModels service via HTTP POST."),
            ("hot-Patching with update_tool", "Allows instant hot-patching of SAS algorithms or UI parameters in active production deployments.")
        ],
        subtitle="Tier 13: Packaging & Deployment",
        mode='wide'
    )

    set_slide_notes(
        s9,
        "Slide 9 highlights the specialized MCP tools we engineered specifically for CAF authoring in Tier 13.\n\n"
        "First, describe_caf_schema. LLM context windows are precious. You don't want to load thousands of lines of XML schema into every prompt. "
        "This tool enables progressive disclosure: the agent asks for the parameter schema, then the lookup schema, then the interaction rules, authoring each piece with precision.\n\n"
        "Second, validate_analysis_type_spec. The Viya backend API will silently accept a zip with an invalid report token, but the user will see a broken report. "
        "Our validator catches these errors locally before uploading, verifying that every token in the report has a matching producer in the input XML.\n\n"
        "Third, templatize_va_report. Rather than hand-crafting complex XML reports, a business analyst designs the chart in Visual Analytics, and this tool converts it into a CAF template automatically. "
        "Then create_analysis_type deploys it straight to Viya."
    )

    # =========================================================================
    # SLIDE 10: Domain 3 (FQA) - Motivation: Field Quality & Warranty Triage
    # =========================================================================
    s10 = prs.slides.add_slide(layout_content)
    set_slide_header(
        s10,
        "Field Quality Analytics (FQA): Early Warning at Enterprise Scale",
        "Mitigating Multi-Million-Dollar Warranty Risks in Heavy Industry, Automotive & Commercial Equipment"
    )

    add_card(
        s10, margin_x_3, card_y, card_w_3, card_h, c_wine,
        "The Warranty Cost Crisis",
        [
            ("Massive Financial Exposure", "Warranty claims and recalls consume 2% to 4% of gross revenue for global automotive and heavy machinery OEMs."),
            ("Direct Bottom-Line Loss", "For a large manufacturer, a delayed recall or slow root-cause containment easily exceeds $50M to $100M in unbudgeted warranty reserves."),
            ("Supplier Recovery Bottlenecks", "Recovering costs from tier-1 suppliers requires indisputable statistical proof that defect spikes correlate to specific production lots."),
            ("Regulatory & Safety Stakes", "Late detection of critical brake, steering, or thermal issues risks regulatory fines and catastrophic safety events.")
        ],
        subtitle="Business Motivation",
        mode='wide'
    )

    add_card(
        s10, margin_x_3 + card_w_3 + gap_3, card_y, card_w_3, card_h, c_amber,
        "The Analytical Firehose",
        [
            ("Millions of Claims", "Quality teams are overwhelmed by hundreds of thousands of monthly warranty records spanning multiple vehicle series and geographies."),
            ("Noisy Diagnostic Codes", "Technicians submit inconsistent DTCs, overlapping labor codes, and varied part numbers for the identical mechanical failure."),
            ("Reporting Lag & Censoring", "Claims take 30 to 90 days to filter from dealerships into the data lake, disguising emerging defects as healthy cohorts."),
            ("Analyst Fatigue", "Engineers spend 80% of their workday building repetitive data selections and filtering out false alarms instead of investigating root causes.")
        ],
        subtitle="The Operational Bottleneck",
        mode='wide'
    )

    add_card(
        s10, margin_x_3 + 2 * (card_w_3 + gap_3), card_y, card_w_3, card_h, c_blue,
        "The Autonomous FQA Agent",
        [
            ("Automated Alert Triage", "Scans Emerging Issues runs continuously, detecting statistically significant production-period and event-month spikes."),
            ("Senior Engineer Heuristics", "Encapsulates decades of reliability engineering expertise into a systematic 4-phase investigation playbook."),
            ("Rapid Root Cause Isolation", "Executes Pareto, Exposure, Weibull, and Decision Tree analyses to isolate responsible component lots within minutes."),
            ("Verified Audit Trail", "Generates executive-ready technical briefs complete with verified statistics and supplier containment boundaries.")
        ],
        subtitle="The Agent-Driven Solution",
        mode='wide'
    )

    set_slide_notes(
        s10,
        "Now we arrive at our third domain: SAS Field Quality Analytics (FQA).\n\n"
        "In industries like automotive and commercial vehicles (for instance, MAN Truck & Bus), warranty management is a multi-million-dollar challenge. "
        "Warranty costs routinely represent 2 to 4 percent of revenue. If a defect in a turbocharger, transmission, or battery module goes undetected for three extra months, "
        "thousands of additional defective vehicles leave the factory floor, multiplying recall costs exponentially.\n\n"
        "The operational problem is not lack of data—it's data overload. Quality engineers receive hundreds of statistical alerts every month, buried under thousands of diagnostic trouble codes and noisy dealership notes. "
        "Furthermore, reporting lag means recent vehicle builds appear healthy simply because claims haven't been filed yet.\n\n"
        "We built a specialized FQA Agent with 14 analytical tools that automates the entire investigation funnel, turning a 6-hour manual investigation into a 3-minute verified triage."
    )

    # =========================================================================
    # SLIDE 11: Domain 3 (FQA) - The 4-Phase Guided Investigation Funnel
    # =========================================================================
    s11 = prs.slides.add_slide(layout_content)
    set_slide_header(
        s11,
        "FQA Guided Investigation Funnel: From Alert to Root Cause",
        "Codified Senior Reliability Engineering Heuristics Enforced by the FQA Agent Skill"
    )

    add_card(
        s11, margin_x_4, card_y, card_w_4, card_h, c_blue,
        "Phase 0: Triage",
        [
            ("Alert Grid Screening", "Evaluates Emerging Issues runs across build-month and event-month matrices."),
            ("Cluster Detection", "Identifies diagonal (calendar-time) vs. horizontal (build-period defect) failure signatures."),
            ("Priority Ranking", "Ranks alerts by unit volume, excess claim rate, and growth velocity to surface high-risk issues.")
        ],
        subtitle="Filter the Noise",
        mode='medium'
    )

    add_card(
        s11, margin_x_4 + card_w_4 + gap_4, card_y, card_w_4, card_h, c_teal,
        "Phase 1: Descriptive",
        [
            ("Pareto Analysis", "Ranks top labor operation codes and replaced part numbers to identify the failing sub-assembly."),
            ("Summary Tables", "Cross-tabulates failure rates across vehicle chassis, engine types, and assembly plants."),
            ("Geographic Heatmaps", "Detects regional clustering to determine if climate, terrain, or fuel quality contributes to failure.")
        ],
        subtitle="Where & What",
        mode='medium'
    )

    add_card(
        s11, margin_x_4 + 2 * (card_w_4 + gap_4), card_y, card_w_4, card_h, c_amber,
        "Phase 4: Normalization",
        [
            ("Exposure Adjustment", "Normalizes claim rates by Time in Service (TIS) at 30, 60, and 90-day intervals, eliminating mileage bias."),
            ("Weibull Shape Beta", "Calculates β: β < 1 indicates infant mortality (manufacturing defect); β > 1 indicates wear-out."),
            ("Build Period Cutoff", "Determines exact production start date where the failure rate diverged from baseline.")
        ],
        subtitle="Normalize Early!",
        mode='medium'
    )

    add_card(
        s11, margin_x_4 + 3 * (card_w_4 + gap_4), card_y, card_w_4, card_h, c_green,
        "Phase 2: Segmentation",
        [
            ("Statistical Drivers", "Identifies variables that exhibit statistically significant divergence between failed and healthy units."),
            ("Decision Tree Splits", "Generates high-purity cohort splits (e.g., Supplier Lot X + Plant Y + Software Version Z)."),
            ("Containment Package", "Outputs exact VIN boundary lists for immediate manufacturing quarantine and supplier chargeback.")
        ],
        subtitle="Isolate Root Cause",
        mode='medium'
    )

    set_slide_notes(
        s11,
        "Slide 11 illustrates the 4-phase investigation funnel codified into the FQA Agent's knowledge base.\n\n"
        "Phase 0 is Triage: The agent scans early warning runs. It distinguishes between diagonal patterns (which indicate external calendar shocks like winter freeze) and horizontal patterns (which indicate defective build lots).\n\n"
        "Phase 1 is Descriptive: Running Pareto analyses on labor and part codes to pinpoint the physical component, followed by geographic and platform breakdowns.\n\n"
        "Phase 4 is Normalization, which our playbook mandates running EARLY: You cannot evaluate raw claim counts because older vehicles naturally accumulate more claims. "
        "The agent normalizes by Time in Service (TIS) and fits a Weibull reliability distribution. If the Weibull shape parameter beta is less than 1, you have an infant mortality defect from the factory floor; if beta is greater than 1, you have a wear-out issue.\n\n"
        "Phase 2 is Segmentation: Running Statistical Drivers and Decision Trees to isolate the precise vehicle configuration, assembly line, and supplier batch responsible for the excess claims."
    )

    # =========================================================================
    # SLIDE 12: Domain 4 (Forecasting) - Motivation & Code-First Strategy
    # =========================================================================
    s12 = prs.slides.add_slide(layout_content)
    set_slide_header(
        s12,
        "CPG Hierarchical Forecasting: Beyond GUI Model Studio",
        "Distributed Hybrid Tournament Pitting SAS Visual Forecasting Against Modern Open-Source Stacks"
    )

    add_card(
        s12, margin_x_3, card_y, card_w_3, card_h, c_blue,
        "The Enterprise Forecasting Challenge",
        [
            ("Walmart M5 Dataset", "Forecasting retail hierarchies: thousands of items nested within departments, categories, stores, and states."),
            ("Intermittent & Lumpy Demand", "High prevalence of zero-demand periods requiring Syntetos-Boylan intermittency categorization (ADI/CV²)."),
            ("Multi-Level Coherence", "Forecasts generated at the item level must mathematically reconcile with store, regional, and national demand totals."),
            ("Supply Chain Volatility", "Stockouts and promotional spikes require robust models capable of handling price elasticity and calendar events.")
        ],
        subtitle="Retail Scale & Complexity",
        mode='wide'
    )

    add_card(
        s12, margin_x_3 + card_w_3 + gap_3, card_y, card_w_3, card_h, c_wine,
        "Breaking Free from the GUI",
        [
            ("The Model Studio Bottleneck", "GUI-based drag-and-drop forecasting tools limit automation, impede version control, and cannot integrate into automated CI/CD pipelines."),
            ("Pure-Code Headless Pipeline", "Orchestrated end-to-end via Python CLI (submit_viya_extlang.py) communicating with the SAS Viya Compute Server API."),
            ("Reproducible Batch Execution", "Fully scriptable parameter tuning, data prep, model fitting, and validation without human manual clicking."),
            ("Automated Champion Selection", "Systematic evaluation over 9 industry metrics (MASE, RMSSE, WAPE, Bias, FA) to crown the winning model per node.")
        ],
        subtitle="Automated Code-First Architecture",
        mode='wide'
    )

    add_card(
        s12, margin_x_3 + 2 * (card_w_3 + gap_3), card_y, card_w_3, card_h, c_green,
        "The Hybrid Model Tournament",
        [
            ("SAS ATSM Engine", "Industrial-grade Automatic Time Series Modeling: automated ARIMA, Exponential Smoothing (ESM), and Unobserved Components (UCM)."),
            ("Nixtla Open-Source Stack", "Benchmarked against state-of-the-art Python algorithms: StatsForecast (AutoARIMA, CES), MLForecast (LightGBM), and NeuralForecast."),
            ("Prophet & Foundation Models", "Cross-benchmarked with Meta's Prophet and zero-shot foundation models like Amazon Chronos."),
            ("Objective Bake-Off", "Pits proprietary SAS algorithms directly against open-source contenders on identical cross-validation folds.")
        ],
        subtitle="SAS Visual Forecasting vs. OSS",
        mode='wide'
    )

    set_slide_notes(
        s12,
        "Now let's examine our fourth domain: Enterprise Hierarchical Forecasting on SAS Viya.\n\n"
        "In supply chain and consumer packaged goods (CPG)—such as our work benchmarking retail forecasting on the famous Walmart M5 dataset—organizations must forecast demand across millions of item-store combinations. "
        "These series exhibit extreme intermittency, seasonal spikes, and hierarchical constraints.\n\n"
        "Traditionally, organizations use SAS Model Studio, which is a wonderful visual UI. But modern engineering teams need automated, headless, code-first pipelines that can be triggered by APIs, scheduled via CI/CD, and orchestrated by AI agents.\n\n"
        "We built a pure-code architecture that runs a massive hybrid forecasting tournament on SAS Viya. "
        "We pit SAS Visual Forecasting (PROC TSMODEL with the ATSM package) directly against the Nixtla open-source ecosystem—StatsForecast, MLForecast, Meta Prophet, and deep learning foundation models like Chronos."
    )

    # =========================================================================
    # SLIDE 13: Domain 4 (Forecasting) - Distributed Execution & K8s Resilience
    # =========================================================================
    s13 = prs.slides.add_slide(layout_content)
    set_slide_header(
        s13,
        "Distributed CAS EXTLANG, MinT Reconciliation & K8s Resilience",
        "Massively Parallel In-Memory Execution under Strict 32Gi Kubernetes Pod Constraints"
    )

    add_card(
        s13, margin_x_3, card_y, card_w_3, card_h, c_teal,
        "CAS EXTLANG Distributed Python",
        [
            ("Parallel Worker Execution", "Uses PROC TSMODEL with EXTLANG to dispatch Python worker scripts directly onto distributed CAS in-memory worker nodes."),
            ("Zero Data Shuffling", "Data partitions remain in CAS RAM; Python processes execute locally against each time series chunk in parallel."),
            ("PROC PYTHON & MP-CONNECT", "Global panel models (NeuralForecast) execute in Compute Server pods using asynchronous SAS MP-CONNECT sessions."),
            ("Hybrid Telemetry Logging", "Captures per-series training and scoring durations via OUTSCALAR to profile exact algorithm computational cost.")
        ],
        subtitle="High-Throughput Execution",
        mode='wide'
    )

    add_card(
        s13, margin_x_3 + card_w_3 + gap_3, card_y, card_w_3, card_h, c_blue,
        "Optimal MinT Reconciliation",
        [
            ("Hierarchical Incoherence", "Independently generated forecasts at item, department, and state levels never add up, creating supply chain chaos."),
            ("Trace Minimization (MinT)", "Implemented optimal forecast shrinkage in PROC IML using sample error covariance matrices."),
            ("Mathematical Optimality", "MinT provably produces unbiased, coherent forecasts with minimal error variance across all hierarchy levels."),
            ("Automated Dashboard Export", "Outputs reconciled hierarchical leaderboards and confidence intervals for visual inspection.")
        ],
        subtitle="PROC IML Optimization",
        mode='wide'
    )

    add_card(
        s13, margin_x_3 + 2 * (card_w_3 + gap_3), card_y, card_w_3, card_h, c_amber,
        "Resilience on Constrained K8s",
        [
            ("Strict Hardware Ceilings", "Viya Compute Server pods running with tight limits: 4 CPUs and a strict 32Gi memory ceiling."),
            ("Bypassing REST Limits", "Viya Compute HTTP APIs fail on files >30MB; solved by uploading compressed .zip payloads and extracting via SAS ZIP engine."),
            ("Binary fcopy Protocol", "Transfers files from filesrvc virtual storage to physical WORK directories using binary fcopy(RECFM=N)."),
            ("OOM Prevention Guardrails", "Runtime telemetry in mycas.runtime_report detects memory spikes before pods trigger Kubernetes OOMKills.")
        ],
        subtitle="Production Engineering",
        mode='wide'
    )

    set_slide_notes(
        s13,
        "Slide 13 reveals the deep engineering required to run distributed forecasting on modern Kubernetes clusters.\n\n"
        "First, CAS EXTLANG. How do you run Python models like Prophet or LightGBM on hundreds of thousands of series? "
        "We used the EXTLANG framework inside PROC TSMODEL. CAS distributes the time series data across its worker nodes, and each node launches Python threads locally in parallel. There is zero network data shuffling.\n\n"
        "Second, MinT reconciliation. If the bottom-level item forecasts don't sum to the store-level forecasts, procurement breaks down. "
        "We implemented Trace Minimization (MinT) in PROC IML, calculating the optimal shrinkage matrix to make the entire retail hierarchy mathematically coherent.\n\n"
        "Third, Kubernetes resilience. In enterprise clusters, compute pods are strictly capped—in our case, 4 CPUs and a 32Gi memory ceiling. "
        "Standard HTTP REST file uploads choke on large datasets. We engineered a streaming architecture using compressed ZIP payloads, SAS binary fcopy, and runtime telemetry that monitors memory per series to prevent cluster crashes."
    )

    # =========================================================================
    # SLIDE 14: Cross-Domain Synergy: The Unified Industrial Copilot
    # =========================================================================
    s14 = prs.slides.add_slide(layout_content)
    set_slide_header(
        s14,
        "The Unified Industrial AI Loop: Chaining SAS Capabilities",
        "How an Autonomous Agent Coordinates Streaming, Custom IoT, Quality, and Supply Chain Forecasting"
    )

    add_card(
        s14, margin_x_4, card_y, card_w_4, card_h, c_blue,
        "1. Real-Time Detection",
        [
            ("ESP Ingestion", "Continuous Kafka telemetry streams through SAS ESP."),
            ("Edge Outlier Trigger", "ESP pattern matching flags abnormal multi-sensor vibration or thermal signature in production unit.")
        ],
        subtitle="Edge Streaming (ESP)",
        mode='medium'
    )

    add_card(
        s14, margin_x_4 + card_w_4 + gap_4, card_y, card_w_4, card_h, c_teal,
        "2. Custom ML Isolation",
        [
            ("CAF Model Execution", "Agent launches SVDD anomaly detection analysis type via MCP Tier 13."),
            ("Multi-Sensor Isolation", "Confirms multivariate boundary deviation and isolates failing sensor tags.")
        ],
        subtitle="Custom IoT (CAF)",
        mode='medium'
    )

    add_card(
        s14, margin_x_4 + 2 * (card_w_4 + gap_4), card_y, card_w_4, card_h, c_amber,
        "3. Warranty Investigation",
        [
            ("FQA Funnel Execution", "FQA Agent checks fleet warranty claims for the flagged component."),
            ("Supplier Containment", "Isolates the defective build lot (Weibull β < 1) and generates supplier recovery brief.")
        ],
        subtitle="Quality Analytics (FQA)",
        mode='medium'
    )

    add_card(
        s14, margin_x_4 + 3 * (card_w_4 + gap_4), card_y, card_w_4, card_h, c_green,
        "4. Supply Adjustment",
        [
            ("Hierarchical ATSM", "Agent re-runs supply chain forecasting tournament via ATSM."),
            ("Spares Reallocation", "Automatically recalculates spare parts demand and updates regional distribution quotas.")
        ],
        subtitle="Forecasting (ATSM)",
        mode='medium'
    )

    set_slide_notes(
        s14,
        "Slide 14 brings our entire story together into a unified operational loop.\n\n"
        "In traditional organizations, these four departments operate in total isolation: the IoT engineering team uses ESP; the plant analytics team uses CAF; the warranty team uses FQA; "
        "and the supply chain team uses Forecasting. Hand-offs take weeks of emails, spreadsheets, and meetings.\n\n"
        "With the SAS MCP Server, an autonomous AI agent acts as the cross-departmental orchestrator:\n"
        "Step 1: ESP flags a continuous streaming vibration anomaly on the assembly line.\n"
        "Step 2: The agent triggers a CAF SVDD model to confirm the multivariate boundary breach.\n"
        "Step 3: The agent cross-references warranty data in FQA, identifying an emerging infant-mortality defect in the same component lot.\n"
        "Step 4: The agent executes our ATSM forecasting tournament to adjust demand forecasts for replacement spare parts across the regional dealer network.\n\n"
        "This is not science fiction—it is the direct result of connecting SAS Viya's analytical engines via MCP."
    )

    # =========================================================================
    # SLIDE 15: Live Demonstration Playbook
    # =========================================================================
    s15 = prs.slides.add_slide(layout_content)
    set_slide_header(
        s15,
        "Live Agent Demonstration: The 4-Act Operational Walkthrough",
        "Real-Time MCP Interaction Demonstrating Autonomous Analytics on Live SAS Viya Backends"
    )

    add_card(
        s15, margin_x_4, card_y, card_w_4, card_h, c_blue,
        "Act 1: Environment Discovery",
        [
            ("Agent Prompt", "'Inspect the current Viya environment, list active CAS tables, and check recent FQA data selections.'"),
            ("Tools Executed", "list_castables, list_data_selections_tool, get_castable_info."),
            ("Result", "Agent returns structured summary of active bioreactor and warranty cohorts in seconds.")
        ],
        subtitle="Zero-Click Discovery",
        mode='medium'
    )

    add_card(
        s15, margin_x_4 + card_w_4 + gap_4, card_y, card_w_4, card_h, c_teal,
        "Act 2: Warranty Triage",
        [
            ("Agent Prompt", "'Check the latest Emerging Issues run and investigate any alert with excess claims > 25%.'"),
            ("Tools Executed", "list_emerging_issue_runs_tool, list_alerts_for_run_tool, run_pareto_analysis_tool."),
            ("Result", "Agent identifies an active defect spike, verifies preflights, and generates an exposure-adjusted Pareto.")
        ],
        subtitle="Automated Triage",
        mode='medium'
    )

    add_card(
        s15, margin_x_4 + 2 * (card_w_4 + gap_4), card_y, card_w_4, card_h, c_amber,
        "Act 3: CAF Deployment",
        [
            ("Agent Prompt", "'Validate the new SVDD anomaly detection package and deploy it to AIoT.'"),
            ("Tools Executed", "describe_caf_schema, validate_analysis_type_spec, create_analysis_type."),
            ("Result", "Agent lints XML tokens, checks BIRD bindings, and registers the custom analytic in Viya.")
        ],
        subtitle="Analysis-as-Code",
        mode='medium'
    )

    add_card(
        s15, margin_x_4 + 3 * (card_w_4 + gap_4), card_y, card_w_4, card_h, c_green,
        "Act 4: Narrative Synthesis",
        [
            ("Agent Output", "Synthesizes complete engineering root-cause brief with exact VIN boundaries and supplier chargeback numbers."),
            ("Evidence", "Every statistic links to verified CAS tables and procedure logs."),
            ("Zero Hallucinations", "Audit-ready report ready for engineering leadership.")
        ],
        subtitle="Verified Briefing",
        mode='medium'
    )

    set_slide_notes(
        s15,
        "Slide 15 serves as our live demonstration transition or playbook.\n\n"
        "We walk through a 4-act live demonstration of the MCP server connected to Claude Code or SAS RAM:\n"
        "In Act 1, we prompt the agent to discover active CAS tables and FQA data selections. The agent calls list_castables and list_data_selections_tool, giving an immediate environmental overview.\n"
        "In Act 2, we ask the agent to triage recent early warning alerts. It executes list_alerts_for_run_tool, performs preflight checks, runs a Pareto analysis, and isolates the failing component.\n"
        "In Act 3, we demonstrate Analysis-as-Code: the agent lints an SVDD anomaly analysis package using validate_analysis_type_spec and deploys it live to the AIoT service.\n"
        "In Act 4, the agent generates an executive briefing where every single number is hyperlinked to verified SAS CAS tables and procedure logs."
    )

    # =========================================================================
    # SLIDE 16: Closing Slide & Strategic Value
    # =========================================================================
    s16 = prs.slides.add_slide(layout_closing)
    for sh in s16.shapes:
        if sh.is_placeholder:
            if sh.placeholder_format.idx == 0:
                p = sh.text_frame.paragraphs[0]
                p.text = "Autonomous Analytics on SAS Viya"
                p.font.name = font_family
                p.font.bold = True
                p.font.size = Pt(44)
            elif sh.placeholder_format.idx == 10:
                p = sh.text_frame.paragraphs[0]
                p.text = "Model Context Protocol (MCP) | SAS RAM | Production Ready"
                p.font.name = font_family
                p.font.size = Pt(20)

    set_slide_notes(
        s16,
        "To close: What we have demonstrated today is that SAS Viya is not just a legacy statistical platform—it is the ultimate high-performance analytical engine for the AI agent era.\n\n"
        "By building the SAS MCP Server with structured concurrency, strict statistical guardrails, and deep integrations across ESP, CAF, FQA, and ATSM, "
        "we have unlocked autonomous, verified industrial analytics.\n\n"
        "The server is containerized, published on Azure Container Registry for SAS RAM, and ready for deployment. "
        "Thank you for your time, and I welcome any questions."
    )

    prs.save(output_path)
    print(f"Presentation B successfully generated at: {output_path}")

if __name__ == '__main__':
    build_presentation_b()
